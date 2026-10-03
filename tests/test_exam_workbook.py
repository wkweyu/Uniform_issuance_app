from io import BytesIO

from openpyxl import load_workbook

from blueprints.exams.workbook import (
    ExamWorkbookError,
    build_component_marks_workbook,
    validate_component_marks_workbook,
)


def _context():
    return {
        12: {
            'subject': {'id': 12, 'name': 'Mathematics', 'code': 'MTH'},
            'components': [
                {
                    'id': 101, 'name': 'CAT', 'maximum_mark': 20,
                    'is_required': True,
                },
                {
                    'id': 102, 'name': 'Final', 'maximum_mark': 80,
                    'is_required': True,
                },
            ],
            'students': [
                {'AdmNo': '1001', 'FName': 'Ada', 'LName': 'Lovelace'},
                {'AdmNo': '1002', 'FName': 'Grace', 'LName': 'Hopper'},
            ],
        },
    }


def test_workbook_template_preserves_roster_and_has_component_columns():
    content = build_component_marks_workbook(list(_context().values()))
    workbook = load_workbook(BytesIO(content), data_only=False)

    assert workbook.sheetnames == ['MTH - 12']
    sheet = workbook.active
    assert sheet.max_row == 3
    assert sheet['A2'].value == '1001'
    assert sheet['C1'].value == 'mark_101'
    assert sheet['D1'].value == 'absent_101'
    assert sheet['I1'].value == 'server_calculated_total'
    workbook.close()


def test_workbook_validation_preserves_zero_and_absence_as_distinct_states():
    content = build_component_marks_workbook(list(_context().values()))
    workbook = load_workbook(BytesIO(content))
    sheet = workbook.active
    sheet['C2'] = 0
    sheet['D2'] = 'FALSE'
    sheet['E2'] = 'Recorded zero'
    sheet['F2'] = 20
    sheet['G2'] = 'FALSE'
    sheet['C3'] = None
    sheet['D3'] = 'TRUE'
    sheet['F3'] = 40
    sheet['G3'] = 'FALSE'
    workbook.active['I2'] = '=SUM(C2,F2)'
    output = BytesIO()
    workbook.save(output)
    workbook.close()

    result = validate_component_marks_workbook(output.getvalue(), _context())

    assert result['valid'] is True
    assert len(result['prepared_marks']) == 4
    zero = next(row for row in result['prepared_marks'] if row['student_id'] == '1001' and row['component_id'] == 101)
    absent = next(row for row in result['prepared_marks'] if row['student_id'] == '1002' and row['component_id'] == 101)
    assert zero['mark'] == 0 and zero['is_absent'] is False
    assert absent['mark'] is None and absent['is_absent'] is True
    assert all(row['mark'] is not None for row in result['prepared_marks'] if row['student_id'] == '1002' and row['component_id'] == 102)


def test_workbook_validation_reports_missing_roster_member_and_conflicting_absence():
    content = build_component_marks_workbook(list(_context().values()))
    workbook = load_workbook(BytesIO(content))
    sheet = workbook.active
    sheet.delete_rows(3)
    sheet['C2'] = 10
    sheet['D2'] = 'TRUE'
    output = BytesIO()
    workbook.save(output)
    workbook.close()

    result = validate_component_marks_workbook(output.getvalue(), _context())

    assert result['valid'] is False
    messages = [error['message'] for error in result['errors']]
    assert any('both absent and scored' in message for message in messages)
    assert any('omits 1 eligible roster learner' in message for message in messages)


def test_workbook_validation_rejects_missing_subject_sheet():
    content = build_component_marks_workbook(list(_context().values()))

    result = validate_component_marks_workbook(
        content,
        {
            **_context(),
            13: {
                **_context()[12],
                'subject': {'id': 13, 'name': 'English', 'code': 'ENG'},
            },
        },
    )

    assert result['valid'] is False
    assert any('13' in error['message'] for error in result['errors'])


def test_workbook_validation_rejects_oversized_content():
    try:
        validate_component_marks_workbook(b'x' * (10 * 1024 * 1024 + 1), _context())
    except ExamWorkbookError as error:
        assert '10 MiB' in str(error)
    else:
        raise AssertionError('oversized workbook was accepted')
