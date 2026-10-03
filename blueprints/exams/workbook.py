"""Bounded XLSX template and validation helpers for component mark imports."""

from io import BytesIO
import re
from decimal import Decimal, InvalidOperation
from typing import Dict, List, Mapping, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill


MAX_WORKBOOK_BYTES = 10 * 1024 * 1024
MAX_WORKBOOK_SHEETS = 100
MAX_WORKBOOK_ROWS = 5000


class ExamWorkbookError(ValueError):
    """Raised when an examination workbook cannot be safely parsed."""


def _safe_sheet_name(subject: Mapping) -> str:
    title = re.sub(r'[\[\]:*?/\\]', '-', str(subject.get('code') or subject['name']))
    title = title.strip(" '")[:22] or 'Subject'
    return f'{title} - {subject["id"]}'


def build_component_marks_workbook(subject_contexts: Sequence[Mapping]) -> bytes:
    """Create one learner-roster worksheet per configured exam subject."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    for context in subject_contexts:
        subject = context['subject']
        components = context['components']
        if not components:
            continue
        sheet = workbook.create_sheet(_safe_sheet_name(subject))
        headers = ['student_id', 'student_name']
        for component in components:
            headers.extend((
                f"mark_{component['id']}",
                f"absent_{component['id']}",
                f"remarks_{component['id']}",
            ))
        headers.extend(('server_calculated_total', 'server_calculated_percentage'))
        sheet.append(headers)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='253858')
        for student in context['students']:
            row = [str(student['AdmNo']), f"{student['FName']} {student['LName']}"]
            for _component in components:
                row.extend(('', 'FALSE', ''))
            row.extend(('', ''))
            sheet.append(row)
        sheet.freeze_panes = 'C2'
        sheet.auto_filter.ref = sheet.dimensions
        sheet.column_dimensions['A'].width = 18
        sheet.column_dimensions['B'].width = 32
        for column in range(3, len(headers) + 1):
            sheet.column_dimensions[
                sheet.cell(row=1, column=column).column_letter
            ].width = 23
    if not workbook.worksheets:
        raise ExamWorkbookError(
            "This class has no configured assessment components to export."
        )
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def validate_component_marks_workbook(
    content: bytes,
    subject_contexts: Mapping[int, Mapping],
) -> Dict:
    """Validate all configured sheets without writing marks to the database."""
    if not content:
        raise ExamWorkbookError("The uploaded workbook is empty.")
    if len(content) > MAX_WORKBOOK_BYTES:
        raise ExamWorkbookError("Workbook exceeds the 10 MiB upload limit.")
    try:
        workbook = load_workbook(
            BytesIO(content),
            read_only=True,
            data_only=False,
        )
    except Exception as exc:
        raise ExamWorkbookError(
            "The upload is not a readable .xlsx workbook."
        ) from exc

    if len(workbook.worksheets) > MAX_WORKBOOK_SHEETS:
        workbook.close()
        raise ExamWorkbookError(
            f"A workbook cannot contain more than {MAX_WORKBOOK_SHEETS} sheets."
        )

    errors = []
    warnings = []
    prepared_marks = []
    seen_sheet_subjects = set()
    expected_roster_by_subject = {}
    try:
        for sheet in workbook.worksheets:
            match = re.search(r'(?:\s-\s|\s)(\d+)$', sheet.title)
            if not match:
                errors.append({
                    'sheet': sheet.title,
                    'row': 1,
                    'message': "Sheet name must end with its subject ID, e.g. MTH - 12.",
                })
                continue
            subject_id = int(match.group(1))
            context = subject_contexts.get(subject_id)
            if context is None:
                errors.append({
                    'sheet': sheet.title,
                    'row': 1,
                    'message': "This subject is not configured for this exam class.",
                })
                continue
            if subject_id in seen_sheet_subjects:
                errors.append({
                    'sheet': sheet.title,
                    'row': 1,
                    'message': "Duplicate worksheet for this subject.",
                })
                continue
            seen_sheet_subjects.add(subject_id)
            components = context['components']
            roster = {
                str(student['AdmNo']): student
                for student in context['students']
            }
            expected_roster_by_subject[subject_id] = set(roster)
            rows = sheet.iter_rows(values_only=True)
            headers = next(rows, None)
            if not headers:
                errors.append({
                    'sheet': sheet.title,
                    'row': 1,
                    'message': "Worksheet is missing its header row.",
                })
                continue
            normalized_headers = [
                str(value or '').strip().casefold() for value in headers
            ]
            if len(normalized_headers) != len(set(normalized_headers)):
                errors.append({
                    'sheet': sheet.title,
                    'row': 1,
                    'message': "Worksheet contains duplicate column headers.",
                })
                continue
            header_indexes = {
                name: index for index, name in enumerate(normalized_headers)
            }
            required_headers = {'student_id'}
            for component in components:
                required_headers.update({
                    f"mark_{component['id']}",
                    f"absent_{component['id']}",
                })
            missing_headers = sorted(required_headers - set(header_indexes))
            if missing_headers:
                errors.append({
                    'sheet': sheet.title,
                    'row': 1,
                    'message': f"Missing columns: {', '.join(missing_headers)}.",
                })
                continue

            seen_students = set()
            sheet_row_count = 0
            for row_number, values in enumerate(rows, start=2):
                if row_number > MAX_WORKBOOK_ROWS + 1:
                    errors.append({
                        'sheet': sheet.title,
                        'row': row_number,
                        'message': f"Worksheet exceeds {MAX_WORKBOOK_ROWS} data rows.",
                    })
                    break
                if not any(value not in (None, '') for value in values):
                    continue
                sheet_row_count += 1
                student_id = str(
                    values[header_indexes['student_id']] or ''
                ).strip()
                if not student_id:
                    errors.append({
                        'sheet': sheet.title,
                        'row': row_number,
                        'message': "Student ID is required.",
                    })
                    continue
                if student_id not in roster:
                    errors.append({
                        'sheet': sheet.title,
                        'row': row_number,
                        'message': f"Unknown or ineligible student ID {student_id}.",
                    })
                    continue
                if student_id in seen_students:
                    errors.append({
                        'sheet': sheet.title,
                        'row': row_number,
                        'message': f"Duplicate student ID {student_id}.",
                    })
                    continue
                seen_students.add(student_id)
                for component in components:
                    component_id = component['id']
                    raw_mark = values[header_indexes[f'mark_{component_id}']]
                    raw_absent = values[header_indexes[f'absent_{component_id}']]
                    absent_text = str(raw_absent or '').strip().casefold()
                    if absent_text not in {'', '0', '1', 'true', 'false', 'yes', 'no', 'y', 'n'}:
                        errors.append({
                            'sheet': sheet.title,
                            'row': row_number,
                            'message': f"Invalid absent flag for component {component_id}.",
                        })
                        continue
                    is_absent = absent_text in {'1', 'true', 'yes', 'y'}
                    if is_absent and raw_mark not in (None, ''):
                        errors.append({
                            'sheet': sheet.title,
                            'row': row_number,
                            'message': f"Component {component_id} cannot be both absent and scored.",
                        })
                        continue
                    if raw_mark in (None, ''):
                        mark = None
                        if not is_absent and component.get('is_required', True):
                            errors.append({
                                'sheet': sheet.title,
                                'row': row_number,
                                'message': f"Required component {component_id} is blank; enter a mark or mark absent.",
                            })
                            continue
                    else:
                        try:
                            mark = Decimal(str(raw_mark))
                        except (InvalidOperation, TypeError, ValueError):
                            errors.append({
                                'sheet': sheet.title,
                                'row': row_number,
                                'message': f"Component {component_id} mark must be numeric.",
                            })
                            continue
                        maximum = Decimal(str(component['maximum_mark']))
                        if not mark.is_finite() or mark < 0 or mark > maximum:
                            errors.append({
                                'sheet': sheet.title,
                                'row': row_number,
                                'message': f"Component {component_id} mark must be between 0 and {maximum}.",
                            })
                            continue
                    remark_key = f'remarks_{component_id}'
                    remark = (
                        str(values[header_indexes[remark_key]] or '')[:500]
                        if remark_key in header_indexes else ''
                    )
                    prepared_marks.append({
                        'student_id': student_id,
                        'subject_id': subject_id,
                        'component_id': component_id,
                        'mark': mark,
                        'is_absent': is_absent,
                        'remarks': remark,
                        'sheet': sheet.title,
                        'row': row_number,
                    })
            missing_students = set(roster) - seen_students
            if missing_students:
                errors.append({
                    'sheet': sheet.title,
                    'row': 0,
                    'message': (
                        f"Worksheet omits {len(missing_students)} eligible roster learner(s). "
                        "Keep every learner row in the workbook."
                    ),
                })

        missing_sheets = set(subject_contexts) - seen_sheet_subjects
        for subject_id in sorted(missing_sheets):
            errors.append({
                'sheet': '',
                'row': 0,
                'message': f"Workbook is missing configured subject worksheet {subject_id}.",
            })
    finally:
        workbook.close()

    return {
        'valid': not errors,
        'errors': errors,
        'warnings': warnings,
        'prepared_marks': prepared_marks,
        'sheet_count': len(seen_sheet_subjects),
        'row_count': sum(
            len(roster) for roster in expected_roster_by_subject.values()
        ),
        'mark_count': len(prepared_marks),
    }
